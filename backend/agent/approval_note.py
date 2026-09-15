import os
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional
from backend.config import settings
from backend.security.audit import log_audit_event

logger = logging.getLogger(__name__)

def generate_approval_docx(
    unit_id: str,
    summary: str,
    findings: List[Dict[str, Any]],
    regulatory_refs: List[str],
    recommendations: List[str],
    output_filename: Optional[str] = None
) -> str:
    """
    Generates a professional formatted Word (.docx) Approval Note.
    Saves document to settings.DOWNLOAD_DIR and returns the absolute file path.
    Falls back to a rich .txt report if python-docx is not available.
    """
    if not output_filename:
        output_filename = f"Approval_Note_{unit_id.replace(' ', '_')}.docx"

    # Ensure output filename ends with .docx
    if not output_filename.lower().endswith(".docx"):
        output_filename = output_filename + ".docx"

    output_path = settings.DOWNLOAD_DIR / output_filename
    os.makedirs(settings.DOWNLOAD_DIR, exist_ok=True)

    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = docx.Document()

        # Page Margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.2)
            section.right_margin = Inches(1.2)

        # Title Banner
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run("SOVEREIGN AI WORKBENCH — APPROVAL NOTE")
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(8, 127, 114)

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run("Confidential Industrial & Regulatory Review Document")
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(99, 117, 117)

        doc.add_paragraph()  # Spacer

        # Metadata Table
        table = doc.add_table(rows=4, cols=2)
        table.style = "Table Grid"
        meta_data = [
            ("Target Unit / Asset:", str(unit_id)),
            ("Date of Review:", datetime.now().strftime("%d %B %Y")),
            ("Document Reference:", f"REF-SOV-{unit_id.replace(' ', '-')}-2026"),
            ("Status:", "DRAFT — REVIEW REQUIRED")
        ]
        for idx, (label, val) in enumerate(meta_data):
            cell_0 = table.cell(idx, 0)
            cell_1 = table.cell(idx, 1)
            cell_0.text = label
            cell_1.text = val
            cell_0.paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph()

        # 1. Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=1)
        h1.runs[0].font.color.rgb = RGBColor(20, 38, 38)
        doc.add_paragraph(summary)

        # 2. Key Findings Table
        h2 = doc.add_heading("2. Key Inspection Findings", level=1)
        h2.runs[0].font.color.rgb = RGBColor(20, 38, 38)

        if findings:
            f_table = doc.add_table(rows=1, cols=4)
            f_table.style = "Table Grid"
            hdr_cells = f_table.rows[0].cells
            for i, hdr in enumerate(["ID", "Finding Description", "Severity", "Regulatory Reference"]):
                hdr_cells[i].text = hdr
                hdr_cells[i].paragraphs[0].runs[0].font.bold = True

            for item in findings:
                row_cells = f_table.add_row().cells
                row_cells[0].text = item.get("id", "F-01")
                row_cells[1].text = item.get("title", "Finding")
                row_cells[2].text = item.get("severity", "Medium")
                row_cells[3].text = item.get("ref", "Safety Regulation 2025")
        else:
            doc.add_paragraph("No specific findings recorded.")

        doc.add_paragraph()

        # 3. Regulatory References
        doc.add_heading("3. Regulatory & SOP References", level=1)
        for ref in (regulatory_refs or ["Safety Regulation 2025 — Section 04"]):
            doc.add_paragraph(f"• {ref}", style="List Bullet")

        doc.add_paragraph()

        # 4. Recommendations
        doc.add_heading("4. Recommendations & Corrective Actions", level=1)
        for i, rec in enumerate(recommendations or ["Review findings with authorized engineer."], 1):
            doc.add_paragraph(f"{i}. {rec}")

        doc.add_paragraph()

        # 5. Sign-off section
        doc.add_heading("5. Approval & Review Authorization", level=1)
        notice_p = doc.add_paragraph()
        notice_run = notice_p.add_run(
            "NOTICE: This approval note was drafted by Sovereign AI Workbench local multi-agent system. "
            "Final operational sign-off must be performed by an authorized certifying engineer."
        )
        notice_run.font.italic = True
        notice_run.font.size = Pt(9)
        notice_run.font.color.rgb = RGBColor(99, 117, 117)

        doc.add_paragraph()
        sig_p = doc.add_paragraph("___________________________                ___________________________")
        sig_p2 = doc.add_paragraph("Lead Inspector Signature                      Supervising Engineer Signature")
        sig_p2.runs[0].font.size = Pt(9)

        doc.save(str(output_path))
        logger.info(f"Generated DOCX approval note at: {output_path}")

        log_audit_event(
            action="GENERATE_APPROVAL_DOCX",
            tools_used="python-docx Engine",
            details=f"Generated document {output_filename} for Unit {unit_id}"
        )

        return str(output_path)

    except ImportError as e:
        logger.warning(f"python-docx not available: {e}. Generating rich text fallback.")
    except Exception as e:
        logger.error(f"Error generating DOCX document: {e}. Generating fallback.")

    # Rich TXT fallback if python-docx is unavailable or fails
    txt_filename = output_filename.replace(".docx", ".txt")
    txt_path = settings.DOWNLOAD_DIR / txt_filename
    try:
        sep = "=" * 60
        lines = [
            sep,
            "   SOVEREIGN AI WORKBENCH — APPROVAL NOTE",
            "   Confidential Industrial & Regulatory Review Document",
            sep,
            f"Target Unit / Asset : {unit_id}",
            f"Date of Review      : {datetime.now().strftime('%d %B %Y')}",
            f"Document Reference  : REF-SOV-{unit_id.replace(' ', '-')}-2026",
            f"Status              : DRAFT — REVIEW REQUIRED",
            "",
            "1. EXECUTIVE SUMMARY",
            "-" * 40,
            summary,
            "",
            "2. KEY INSPECTION FINDINGS",
            "-" * 40,
        ]
        for item in (findings or []):
            lines.append(f"  [{item.get('severity','?')}] {item.get('title','Finding')} — Ref: {item.get('ref','')}")
        lines += [
            "",
            "3. REGULATORY & SOP REFERENCES",
            "-" * 40,
        ]
        for ref in (regulatory_refs or []):
            lines.append(f"  • {ref}")
        lines += [
            "",
            "4. RECOMMENDATIONS & CORRECTIVE ACTIONS",
            "-" * 40,
        ]
        for i, rec in enumerate(recommendations or [], 1):
            lines.append(f"  {i}. {rec}")
        lines += [
            "",
            "5. APPROVAL & REVIEW AUTHORIZATION",
            "-" * 40,
            "NOTICE: This approval note was drafted by Sovereign AI Workbench.",
            "Final sign-off must be performed by an authorized certifying engineer.",
            "",
            "Lead Inspector Signature: ___________________________",
            "Supervising Engineer   : ___________________________",
            sep,
        ]
        with open(str(txt_path), "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        logger.info(f"Generated TXT fallback approval note at: {txt_path}")
        log_audit_event(
            action="GENERATE_APPROVAL_TXT",
            tools_used="TXT Fallback Engine",
            details=f"Generated {txt_filename} for Unit {unit_id}"
        )
        return str(txt_path)
    except Exception as e2:
        logger.error(f"Even TXT fallback failed: {e2}")
        return ""
